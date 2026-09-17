# -*- coding: utf-8 -*-
"""
架上寻书 v3 — FastAPI 后端
识别接口使用 SSE 逐块推送进度。
V3 新增：用户注册登录、记录按用户隔离、识别前裁剪、多书搜索、日志落盘。
"""
import base64
import io
import json
import logging
import os
import queue
import re
import secrets
import shutil
import threading
import time
from collections import deque
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from logging.handlers import RotatingFileHandler
from pathlib import Path

from fastapi import Body, Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
import requests
from PIL import Image, ImageOps

from core import auth
from core import scan_login
from core.fuse import fuse_book_lists
from core.merge import merge_fragments
from core.parse import norm_title, parse_books
from core.providers import DEFAULT_PROVIDER, PROVIDER_DEFS, call_vision_any
from core.refine import (assign_shelf_per_row, drop_edge_rows,
                         majority_shelf, snap_rows)
from core.search import search_records
from core.tiles import _y_overlap_min, dedup, dedup_row_x, make_tiles
from core.vision import call_vision

BASE_DIR = Path(__file__).resolve().parent
RESULTS_DIR = BASE_DIR / "results"
RESULTS_DIR.mkdir(exist_ok=True)
LOG_DIR = BASE_DIR / "logs"
LOG_DIR.mkdir(exist_ok=True)

logging.basicConfig(level=logging.INFO)
_file_handler = RotatingFileHandler(
    LOG_DIR / "server.log", maxBytes=2_000_000, backupCount=3, encoding="utf-8")
_file_handler.setFormatter(logging.Formatter(
    "%(asctime)s %(levelname)s %(name)s: %(message)s"))
logging.getLogger().addHandler(_file_handler)
logger = logging.getLogger("bookfinder")

# ---- 访问日志脱敏：?token=xxx 会进入 uvicorn/httpx 的访问日志，落盘前抹掉 ----
_TOKEN_RE = re.compile(r"([?&]token=)[^&\s]+", re.IGNORECASE)


class _TokenRedactor(logging.Filter):
    """把日志记录里 URL 上的 ?token= 查询参数替换为 ***。"""

    def filter(self, record: logging.LogRecord) -> bool:
        try:
            msg = record.getMessage()
        except Exception:
            return True
        if "token=" in msg:
            record.msg = _TOKEN_RE.sub(r"\1***", msg)
            record.args = ()
        return True


# 注意：挂在 logger 上的 filter 只对"直接打到该 logger"的记录生效，
# 子 logger（uvicorn.access / httpx 等）冒泡上来的记录只过 handler 上的 filter，
# 所以 logger 和 handler 都要挂。
_REDACTOR = _TokenRedactor()
logging.getLogger().addFilter(_REDACTOR)
logging.getLogger("uvicorn.access").addFilter(_REDACTOR)
logging.getLogger("uvicorn").addFilter(_REDACTOR)
_file_handler.addFilter(_REDACTOR)
for _h in logging.getLogger().handlers:
    _h.addFilter(_REDACTOR)

CONFIG_PATH = BASE_DIR / "config.json"
CONFIG: dict = {}
if CONFIG_PATH.exists():
    try:
        CONFIG = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    except Exception:
        logger.error("config.json 解析失败，已按空配置启动", exc_info=True)
        CONFIG = {}
if not isinstance(CONFIG, dict):
    CONFIG = {}

# 安全约定：密钥只从环境变量读取，config.json 只保留非敏感配置。
# 历史版本曾把 api_key / secret 明文写进 config.json，这里做兜底：
# 即使文件里还残留这些字段也一律忽略，避免旧配置被重新启用。
for _legacy_key in ("api_key", "secret", "wechat_secret"):
    if _legacy_key in CONFIG:
        CONFIG.pop(_legacy_key)
        logger.warning("config.json 中的 %s 字段已废弃并被忽略，请改用环境变量配置",
                       _legacy_key)

API_KEY = (os.environ.get("DEEPSEEK_API_KEY") or "").strip()
BASE_URL = (os.environ.get("DEEPSEEK_BASE_URL") or CONFIG.get("base_url")
            or "https://api.deepseek.com").rstrip("/")
MODEL = os.environ.get("DEEPSEEK_MODEL") or CONFIG.get("model") or "deepseek-flash"

# 令牌签名密钥：只从环境变量读取，不再回写 config.json。
# 未设置时生成一次性的进程内随机密钥：不落盘（只读文件系统、多实例部署都不会崩），
# 代价是每次重启后所有已签发令牌失效，需要重新登录。
SECRET = (os.environ.get("BOOKFINDER_SECRET") or "").strip()
if not SECRET:
    SECRET = secrets.token_hex(32)
    logger.warning(
        "未设置环境变量 BOOKFINDER_SECRET，已生成临时进程内令牌密钥；"
        "服务重启后所有用户需重新登录。生产环境请固定该环境变量。")
if not API_KEY:
    logger.warning(
        "未设置环境变量 DEEPSEEK_API_KEY，识别接口将直接报错。"
        "启动前请先配置，例如：set DEEPSEEK_API_KEY=sk-xxxxxxxx")

# 多视觉供应商：环境变量里配了 key 的才视为可用（V3 只从环境变量读密钥）
def _available_providers() -> dict[str, dict]:
    out = {}
    for name, d in PROVIDER_DEFS.items():
        key = (os.environ.get(d["env"]) or "").strip()
        # DeepSeek 必须由 DEEPSEEK_API_KEY 提供；其余按各自 env 判定
        out[name] = {
            "label": d["label"], "model": d["model"], "note": d.get("note", ""),
            "api_key": key, "available": bool(key),
        }
    return out
PROVIDERS = _available_providers()

# 微信小程序登录凭据（code2session），未配置则 /api/auth/wechat 返回 503
WECHAT_APPID = os.environ.get("WECHAT_APPID") or CONFIG.get("wechat_appid", "")
WECHAT_SECRET = os.environ.get("WECHAT_SECRET") or ""

auth.init_db()
scan_login.init_db()

MAX_EDGE = 6000
MAX_UPLOAD_BYTES = 30 * 1024 * 1024   # 单次上传体积上限（含 multipart 开销）
MAX_IMAGE_PIXELS = 50_000_000         # 解码前像素上限，防"解压炸弹"
MAX_QUERY_LEN = 100                   # 单个搜索书名最大长度
MAX_EXPORT_IDS = 200                  # 单次导出最多记录数
MAX_NAME_LEN = 200                    # 记录名最大长度
# Pillow 兜底：像素数超过该值 2 倍时直接 DecompressionBombError
Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS
RID_RE = re.compile(r"^\d{8}_\d{6}_\d{3}$")

# ---- 轻量内存级限流（标准库实现，无新依赖）：滑动窗口，进程重启即清空 ----
_RATE_LOCK = threading.Lock()
_RATE_HITS: dict[str, deque[float]] = {}
_RATE_MAX_KEYS = 5000


def _client_ip(request: Request) -> str:
    """取客户端 IP。只信任 socket 对端，不读可被伪造的 X-Forwarded-For。"""
    return (request.client.host if request.client else "") or "unknown"


def _rate_limit(key: str, limit: int, window: float = 60.0) -> None:
    """同一 key 在 window 秒内最多 limit 次，超出抛 429。"""
    now = time.monotonic()
    with _RATE_LOCK:
        hits = _RATE_HITS.get(key)
        if hits is None:
            hits = _RATE_HITS[key] = deque()
        while hits and now - hits[0] > window:
            hits.popleft()
        if len(hits) >= limit:
            raise HTTPException(429, "请求过于频繁，请稍后再试")
        hits.append(now)
        if len(_RATE_HITS) > _RATE_MAX_KEYS:  # 防止 key 无限增长
            for k in [k for k, v in _RATE_HITS.items()
                      if not v or now - v[-1] > window]:
                _RATE_HITS.pop(k, None)


app = FastAPI(title="架上寻书", version="3.0",
              docs_url=None if os.environ.get("BOOKFINDER_ENV") == "prod" else "/docs",
              redoc_url=None if os.environ.get("BOOKFINDER_ENV") == "prod" else "/redoc",
              openapi_url=None if os.environ.get("BOOKFINDER_ENV") == "prod"
              else "/openapi.json")
# 生产模式（单端口托管 dist）下同源，无需放开 CORS；
# 这里只放行 Vite 开发服务器，且 allow_origins 为显式白名单（没有 "*" + credentials）。
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.middleware("http")
async def _guard(request: Request, call_next):
    """上传体积预检（在读取请求体之前拦截）+ 基础安全响应头。"""
    if request.url.path == "/api/recognize":
        cl = request.headers.get("content-length")
        if cl and cl.isdigit() and int(cl) > MAX_UPLOAD_BYTES:
            return JSONResponse(
                {"detail": f"图片过大，请压缩到 {MAX_UPLOAD_BYTES // 1024 // 1024}MB 以内"},
                status_code=413)
    resp = await call_next(request)
    resp.headers.setdefault("X-Content-Type-Options", "nosniff")
    resp.headers.setdefault("X-Frame-Options", "SAMEORIGIN")
    resp.headers.setdefault("Referrer-Policy", "no-referrer")
    if request.url.path.startswith(("/api/", "/results/")):
        # /results/ 的 URL 上带 ?token=，禁止任何中间缓存保存
        resp.headers["Cache-Control"] = "no-store"
    return resp


TOKEN_QUERY_PATHS = ("/results/",)


def require_user(request: Request) -> str:
    """FastAPI 依赖：校验 Bearer 头；?token= 仅对图片路由放行（供 <img> 用）。

    ?token= 会让令牌进入浏览器历史、Referer 与访问日志，因此只保留给无法携带
    请求头的 <img>/小程序 image 组件（/results/ 前缀），其它接口一律只认请求头。
    """
    token = None
    header = request.headers.get("authorization", "")
    if header.startswith("Bearer "):
        token = header[7:]
    elif request.url.path.startswith(TOKEN_QUERY_PATHS):
        token = request.query_params.get("token") or None
    username = auth.verify_token(token, SECRET) if token else None
    if not username:
        raise HTTPException(401, "未登录或登录已过期")
    return username


def _sse(event: str, data: dict) -> str:
    return f"event: {event}\ndata: {json.dumps(data, ensure_ascii=False)}\n\n"


def _new_result_dir() -> tuple[str, Path]:
    for _ in range(10):
        # 只取一次时间：否则秒位与毫秒位可能来自不同的秒（跨秒时 rid 会"回退"，
        # 导致按字典序排序的记录列表出现新的排在旧的前面）
        now = time.time()
        rid = (time.strftime("%Y%m%d_%H%M%S", time.localtime(now))
               + f"_{int(now * 1000) % 1000:03d}")
        rdir = RESULTS_DIR / rid
        try:
            rdir.mkdir()
            return rid, rdir
        except FileExistsError:
            time.sleep(0.005)
    raise RuntimeError("无法创建结果目录")


def _load_result(rid: str) -> dict | None:
    if not RID_RE.match(rid):
        return None
    p = RESULTS_DIR / rid / "result.json"
    if not p.exists():
        return None
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return None


def _public(d: dict) -> dict:
    """返回给客户端的记录：去掉内部的用户归属字段。"""
    return {k: v for k, v in d.items() if k != "user_id"}


def _my_records(user: str) -> list[dict]:
    """当前用户的全部记录，新的在前。"""
    records = []
    for p in sorted(RESULTS_DIR.glob("*/result.json"), reverse=True):
        try:
            d = json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            continue
        if isinstance(d, dict) and d.get("user_id") == user:
            records.append(d)
    return records


def _write_json_atomic(path: Path, data: dict) -> None:
    """临时文件 + os.replace 原子写，避免并发读时读到半个文件。"""
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2),
                   encoding="utf-8")
    os.replace(tmp, path)


_NAME_CTRL_RE = re.compile(r"[\x00-\x1f\x7f]")


def _sanitize_name(s: str) -> str:
    """清洗记录名：去控制字符、只取路径最后一段、截断，避免脏数据进 Excel/前端。"""
    s = _NAME_CTRL_RE.sub("", str(s or "")).replace("\\", "/")
    s = s.split("/")[-1].strip()
    if not s or s in (".", ".."):
        s = "image.jpg"
    return s[:MAX_NAME_LEN]


def _parse_crop(s: str, w: int, h: int) -> tuple[int, int, int, int]:
    """解析归一化裁剪框 "x0,y0,x1,y1"（0~1）为像素坐标，非法抛 400。"""
    try:
        parts = [float(v) for v in s.split(",")]
    except ValueError:
        raise HTTPException(400, "crop 格式错误，应为 x0,y0,x1,y1（0~1）")
    if len(parts) != 4:
        raise HTTPException(400, "crop 格式错误，应为 x0,y0,x1,y1（0~1）")
    x0, y0, x1, y1 = parts
    if not (0 <= x0 < x1 <= 1 and 0 <= y0 < y1 <= 1):
        raise HTTPException(400, "crop 越界或方向错误")
    if (x1 - x0) * (y1 - y0) < 0.01:
        raise HTTPException(400, "crop 面积过小（需 >= 1%）")
    px = (round(x0 * w), round(y0 * h), round(x1 * w), round(y1 * h))
    if px[2] <= px[0] or px[3] <= px[1]:
        # 归一化面积达标但取整后宽/高为 0（极小图），裁剪会抛 ValueError
        raise HTTPException(400, "crop 取整后像素尺寸为 0，请扩大裁剪区域")
    return px


def _fill_from_frame(detail_books: list[dict],
                     frame_books: list[dict]) -> list[dict]:
    """整图框架补漏：某本整图书的 bbox（外扩 10%）内没有任何切块书的中心点，
    说明切块漏了它——以低可信（confidence×0.8，source="frame"）补进结果。

    例外：整图遍的 bbox 定位不准（实测会把同一排的书整体右移 0.17），
    若同一 y 带（y 重叠>0.3）里已有同名（归一化相等或互为子串）的切块书，
    说明切块遍已读到这本物理书，补进去只会以错位框重复，跳过。"""
    centers = [((b["bbox"][0] + b["bbox"][2]) / 2,
                (b["bbox"][1] + b["bbox"][3]) / 2) for b in detail_books]
    filled = []
    for fb in frame_books:
        x1, y1, x2, y2 = fb["bbox"]
        mx, my = (x2 - x1) * 0.1, (y2 - y1) * 0.1
        ex1, ey1, ex2, ey2 = x1 - mx, y1 - my, x2 + mx, y2 + my
        if any(ex1 <= cx <= ex2 and ey1 <= cy <= ey2 for cx, cy in centers):
            continue
        nf = norm_title(fb["title"])
        if nf and any(
                norm_title(db["title"])
                and (norm_title(db["title"]) == nf
                     or norm_title(db["title"]) in nf
                     or nf in norm_title(db["title"]))
                and _y_overlap_min(fb["bbox"], db["bbox"]) > 0.3
                for db in detail_books):
            continue
        nb = dict(fb)
        nb["confidence"] = round(nb.get("confidence", 0.5) * 0.8, 3)
        nb["source"] = "frame"
        filled.append(nb)
    return filled


def _recognize_stream(img: Image.Image, use_tiles: bool):
    """生成器：两遍识别（整图框架 + 切块细节），逐块产出 SSE 事件。

    tiles=1 且确有多个切块时，追加一个整图"框架"调用（与其他切块并行进同一个
    线程池，不增加总耗时）。框架结果不直接进书籍列表，只用于：验证本数差异
    （stats）、补漏（_fill_from_frame）、架标投票。
    """
    w, h = img.size
    tiles = make_tiles(w, h) if use_tiles else [(0, 0, w, h)]
    has_frame = use_tiles and len(tiles) > 1
    tasks = [("detail", t) for t in tiles]
    if has_frame:
        tasks.append(("frame", (0, 0, w, h)))
    total = len(tasks)
    yield _sse("meta", {"tiles_total": total, "width": w, "height": h})

    # 有界队列 + stop 事件：客户端中途断开时工作线程能感知并退出，
    # 不会往无界队列里无限堆积结果（内存增长 + 继续烧 API 额度）。
    events: queue.Queue = queue.Queue(maxsize=max(8, total))
    stop = threading.Event()

    def work(task):
        kind, tile = task
        try:
            x0, y0, x1, y1 = tile
            buf = io.BytesIO()
            img.crop(tile).save(buf, format="JPEG", quality=90)
            raw = call_vision(buf.getvalue(), api_key=API_KEY,
                              base_url=BASE_URL, model=MODEL)
            books = parse_books(raw)
            tw, th = x1 - x0, y1 - y0
            for b in books:
                bx = b["bbox"]
                b["bbox"] = [
                    round((x0 + bx[0] * tw) / w, 4),
                    round((y0 + bx[1] * th) / h, 4),
                    round((x0 + bx[2] * tw) / w, 4),
                    round((y0 + bx[3] * th) / h, 4),
                ]
            return {"pass": kind, "tile": tile, "books": books, "error": None}
        except Exception as e:
            # try 覆盖整个函数体：crop/save 等任何一步失败都返回错误结果，
            # 保证线程池里的异常不会外溢，结束哨兵必然到达
            logger.warning("切块 %s (%s) 识别失败: %s", tile, kind, e)
            return {"pass": kind, "tile": tile, "books": [],
                    "error": str(e)[:200]}

    def _put(item) -> None:
        """带超时的入队，期间持续检查 stop，避免客户端断开后永久阻塞。"""
        while not stop.is_set():
            try:
                events.put(item, timeout=0.5)
                return
            except queue.Full:
                continue

    def run_all():
        pool = ThreadPoolExecutor(max_workers=4)
        try:
            # 用 as_completed 而非 pool.map：map 按提交顺序产出，某一个慢切块会
            # 卡住后面已完成切块的进度推送；as_completed 是"谁先好推谁"。
            futs = [pool.submit(work, t) for t in tasks]
            for fut in as_completed(futs):
                if stop.is_set():
                    break
                if fut.cancelled():
                    continue
                try:
                    res = fut.result()
                except BaseException as e:  # noqa: BLE001 - 兜底，不能让哨兵丢失
                    res = {"pass": "detail", "tile": (0, 0, 0, 0),
                           "books": [], "error": str(e)}
                _put(res)
        except BaseException as e:  # noqa: BLE001
            logger.warning("切块调度异常: %s", e)
        finally:
            pool.shutdown(wait=False, cancel_futures=True)
            _put(None)  # 结束哨兵：finally 保证必然送达，防生成器永久阻塞

    threading.Thread(target=run_all, daemon=True).start()

    detail_books, frame_books, debug = [], [], []
    try:
        while True:
            res = events.get()
            if res is None:
                break
            if res["pass"] == "frame":
                frame_books.extend(res["books"])
            else:
                detail_books.extend(res["books"])
            debug.append({"tile": list(res["tile"]), "pass": res["pass"],
                          "count": len(res["books"]),
                          **({"error": res["error"]} if res["error"] else {})})
            yield _sse("tile", {"tile": list(res["tile"]),
                                "pass": res["pass"],
                                "count": len(res["books"]),
                                "error": res["error"],
                                "done": len(debug), "total": total})
    finally:
        # 生成器被关闭（客户端断开 / GC）时通知工作线程收工
        stop.set()

    detail_debug = [d for d in debug if d["pass"] == "detail"]
    if not detail_books and any(d.get("error") for d in detail_debug) \
            and len(detail_debug) == len(tiles):
        yield _sse("error", {"error": f"所有切块识别均失败: "
                                      f"{detail_debug[0].get('error')}"})
        return

    # 补漏的书在 merge 之前加入；顺序：dedup → merge_fragments →
    # dedup_row_x（物理行内去重）→ drop_edge_rows → snap_rows
    books, shelf_no, stats = _assemble(detail_books, frame_books)
    for i, b in enumerate(books, 1):
        b["index"] = i
    yield _sse("books", {"books": books, "tiles": debug, "shelf": shelf_no,
                         "stats": stats})


def _assemble(detail_books: list[dict], frame_books: list[dict]):
    """把"切块/整图细节遍 + 整图框架遍"的原始书籍组装成最终书籍列表。

    返回 (books, shelf_no, stats)。整图模式（只传 detail，frame 为空）同样走这套
    后处理；多供应商重识别里每家也各自调它，得到"单家已组装"的结果再交由融合。
    """
    has_frame = bool(frame_books)
    filled = _fill_from_frame(detail_books, frame_books) if has_frame else []
    books = merge_fragments(dedup(detail_books + filled))
    n_dropped_row = len(books)
    books = dedup_row_x(books)
    dropped_dup_books = n_dropped_row - len(books)
    if dropped_dup_books:
        logger.info("行内去重丢弃 %d 本（双读幻觉/弱名/碎片）", dropped_dup_books)
    books, dropped_rows, dropped_books = drop_edge_rows(books)
    if dropped_rows:
        logger.info("丢弃边缘半层: %d 行 %d 本", dropped_rows, dropped_books)
    books = snap_rows(books)
    # 架标：先按行（y 带）把每排书的书架号统一为该行的众数完整架标
    # （避免逐本碎片化；一张图两排不同架标各自聚合、不混排），
    # 再做整图级架标：切块书（含补漏）的 shelf 优先投票。整图框架遍里
    # 架标小字分辨率低、误读率高，实测混合投票会把 3/9 张图的分类号带错
    # （见 tests/fragment_benchmark_report.md），因此整图书只在切块完全
    # 没读到架标时作为兜底。
    assign_shelf_per_row(books)
    shelf_no = majority_shelf(books) or majority_shelf(frame_books)
    stats = {"frame_count": len(frame_books),
             "detail_count": len(detail_books), "filled": len(filled),
             "dropped_dup_books": dropped_dup_books,
             "dropped_edge_rows": dropped_rows,
             "dropped_edge_books": dropped_books}
    logger.info("两遍验证: 整图 %d 本, 切块 %d 本, 补漏 %d 本, 合并后 %d 本",
                stats["frame_count"], stats["detail_count"], stats["filled"],
                len(books))
    return books, shelf_no, stats


def _recognize_by_call(img: Image.Image, use_tiles: bool, call_fn):
    """非流式：用 call_fn 对整图（及可选切块）做识别，返回 (detail_books, frame_books)。

    供"重识别特定图"使用：每个视觉供应商各自调一遍，得到单家原始书籍再交给
    _assemble 组装，最后 fuse_book_lists 融合。单个供应商内部任一切块失败只影响该家。
    """
    w, h = img.size
    tiles = make_tiles(w, h) if use_tiles else [(0, 0, w, h)]
    has_frame = use_tiles and len(tiles) > 1

    def work(tile):
        try:
            x0, y0, x1, y1 = tile
            buf = io.BytesIO()
            img.crop(tile).save(buf, format="JPEG", quality=90)
            raw = call_fn(buf.getvalue())
            books = parse_books(raw)
            tw, th = x1 - x0, y1 - y0
            for b in books:
                bx = b["bbox"]
                b["bbox"] = [
                    round((x0 + bx[0] * tw) / w, 4),
                    round((y0 + bx[1] * th) / h, 4),
                    round((x0 + bx[2] * tw) / w, 4),
                    round((y0 + bx[3] * th) / h, 4),
                ]
            return books
        except Exception as e:
            logger.warning("重识别切块 %s 失败: %s", tile, e)
            return []

    detail_books = []
    with ThreadPoolExecutor(max_workers=4) as pool:
        for books in pool.map(work, tiles):
            detail_books.extend(books)

    frame_books = []
    if has_frame:
        try:
            buf = io.BytesIO()
            img.save(buf, format="JPEG", quality=90)
            frame_books = parse_books(call_fn(buf.getvalue()))
        except Exception as e:
            logger.warning("重识别整图框架失败: %s", e)
            frame_books = []
    return detail_books, frame_books


def _provider_call(name: str):
    """构造某供应商的整图调用函数。"""
    p = PROVIDERS[name]
    cfg = PROVIDER_DEFS[name]
    if not p["api_key"]:
        raise RuntimeError(f"未配置 {cfg['label']} 的 API Key")
    def _call(b):
        return call_vision_any(b, base_url=cfg["base_url"], api_key=p["api_key"],
                               model=cfg["model"],
                               thinking=cfg.get("thinking", False),
                               label=cfg["label"])
    _call.__name__ = f"call_{name}"
    return _call


def _fix_name(s: str) -> str:
    """修正非浏览器客户端（如 curl/脚本）以 GBK 字节上传导致的中文文件名乱码。"""
    try:
        raw = s.encode("latin-1")
    except UnicodeEncodeError:
        return s  # 正常 UTF-8 文件名
    for enc in ("utf-8", "gbk"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return s


@app.post("/api/register")
def api_register(request: Request, payload: dict = Body(...)):
    _rate_limit(f"register:ip:{_client_ip(request)}", 60)  # 防批量注册/撞库
    username = str(payload.get("username") or "").strip()
    password = str(payload.get("password") or "")
    if not auth.USERNAME_RE.match(username):
        raise HTTPException(400, "用户名需为2~32位字母/数字/下划线/中文")
    if len(password) < 6:
        raise HTTPException(400, "密码至少6位")
    if len(password) > 128:
        raise HTTPException(400, "密码最长128位")
    if not auth.create_user(username, password):
        raise HTTPException(409, "用户名已存在")
    return {"token": auth.make_token(username, SECRET), "username": username}


@app.post("/api/login")
def api_login(request: Request, payload: dict = Body(...)):
    username = str(payload.get("username") or "").strip()
    password = str(payload.get("password") or "")
    # 双限流：单 IP 总量 + 单 IP×账号 的爆破防护
    _rate_limit(f"login:ip:{_client_ip(request)}", 60)
    _rate_limit(f"login:acct:{_client_ip(request)}:{username[:32]}", 10)
    if len(password) > 128:
        raise HTTPException(400, "密码最长128位")
    if not auth.check_user(username, password):
        raise HTTPException(401, "用户名或密码错误")
    return {"token": auth.make_token(username, SECRET), "username": username}


@app.post("/api/auth/wechat")
def api_auth_wechat(request: Request, payload: dict = Body(...)):
    """微信小程序登录：wx.login() 的 code 调 code2session 换 openid。

    按 openid 查用户，不存在则创建无密码用户（username 为 wx_+openid前8位）。
    """
    if not WECHAT_APPID or not WECHAT_SECRET:
        raise HTTPException(503, "微信登录未配置")
    _rate_limit(f"wxlogin:ip:{_client_ip(request)}", 60)
    code = str(payload.get("code") or "").strip()
    if not code:
        raise HTTPException(400, "缺少 code")
    try:
        resp = requests.get(
            "https://api.weixin.qq.com/sns/jscode2session",
            params={"appid": WECHAT_APPID, "secret": WECHAT_SECRET,
                    "js_code": code, "grant_type": "authorization_code"},
            timeout=10)
    except requests.RequestException as e:
        # 不把原始异常（可能含内网代理地址等）回给客户端，细节只落日志
        logger.warning("微信 code2session 请求失败: %s", e)
        raise HTTPException(502, "微信登录接口不可达，请稍后重试")
    try:
        data = resp.json()
    except ValueError:
        logger.warning("微信 code2session 返回非 JSON")
        raise HTTPException(502, "微信登录接口返回异常，请稍后重试")
    if data.get("errcode"):
        raise HTTPException(401, f"微信登录失败: {data.get('errmsg', '')}")
    openid = data.get("openid")
    if not openid:
        raise HTTPException(401, "微信登录失败: 未返回 openid")
    username = auth.find_user_by_openid(openid) \
        or auth.create_wechat_user(openid)
    return {"token": auth.make_token(username, SECRET), "username": username}


@app.post("/api/auth/wechat/qr")
def api_wechat_qr(request: Request):
    """生成扫码登录票据 + 小程序码。二维码图上带 ticket（scene 参数），
    浏览器轮询 /api/auth/wechat/qr/status 等待确认。"""
    if not (WECHAT_APPID and WECHAT_SECRET) and not scan_login.MOCK_ENABLED:
        raise HTTPException(503, "微信登录未配置")
    _rate_limit(f"wxqr:ip:{_client_ip(request)}", 60)
    ticket = scan_login.create_ticket()
    if scan_login.MOCK_ENABLED:
        # MOCK 模式：占位二维码（SVG data URL），confirm 接口接受任意 code
        return {"ticket": ticket, "qr_image": scan_login.mock_qr_data_url(),
                "expires_in": scan_login.TICKET_TTL}
    try:
        png = scan_login.get_wxa_code_unlimited(WECHAT_APPID, WECHAT_SECRET,
                                                scene=ticket)
    except Exception as e:
        logger.warning("生成小程序码失败: %s", e)
        raise HTTPException(502, "生成小程序码失败，请稍后重试")
    return {"ticket": ticket,
            "qr_image": "data:image/png;base64," + base64.b64encode(png).decode(),
            "expires_in": scan_login.TICKET_TTL}


@app.get("/api/auth/wechat/qr/status")
def api_wechat_qr_status(request: Request, ticket: str = ""):
    """浏览器轮询扫码状态。confirmed 时抢占式消费票据并签发 Bearer token。

    不区分"票不存在"与"已过期"：否则这个接口就成了"某张票是否存在"的探针。
    """
    _rate_limit(f"wxstatus:ip:{_client_ip(request)}", 120)
    ticket = ticket.strip()
    if not scan_login.TICKET_RE.match(ticket):
        raise HTTPException(400, "票据格式不正确")
    rec = scan_login.get_ticket(ticket)
    if rec is None or rec["expires_at"] <= time.time():
        return {"status": "expired"}
    if rec["status"] == "cancelled":
        return {"status": "cancelled"}
    if rec["status"] == "consumed":
        # 已被另一次轮询领取：不暴露"已消费"这个内部状态，按过期处理
        return {"status": "expired"}
    if rec["status"] != "confirmed":
        return {"status": "pending"}
    username = scan_login.claim_confirmed_ticket(ticket)
    if not username:
        # 已被另一个轮询请求消费
        return {"status": "expired"}
    return {"status": "confirmed", "username": username,
            "token": auth.make_token(username, SECRET)}


@app.post("/api/auth/wechat/confirm")
def api_wechat_confirm(request: Request, payload: dict = Body(...)):
    """小程序端确认登录：{ticket, code} → code2session 换 openid → 置 confirmed。

    不返回 token：真正要登录的是电脑浏览器，token 由 qr/status 轮询领取。
    """
    if not (WECHAT_APPID and WECHAT_SECRET) and not scan_login.MOCK_ENABLED:
        raise HTTPException(503, "微信登录未配置")
    _rate_limit(f"wxconfirm:ip:{_client_ip(request)}", 60)
    ticket = str(payload.get("ticket") or "").strip()
    code = str(payload.get("code") or "").strip()
    if not scan_login.TICKET_RE.match(ticket):
        raise HTTPException(400, "票据格式不正确")
    if not code:
        raise HTTPException(400, "缺少 code")
    rec = scan_login.get_ticket(ticket)
    if rec is None or rec["expires_at"] <= time.time():
        raise HTTPException(404, "二维码已过期，请在电脑上刷新后重新扫码")
    if rec["status"] == "cancelled":
        raise HTTPException(409, "该二维码已被取消，请在电脑上重新生成")
    # js_code 一次性有效，失败不能重试；重新扫码会拿到新 code，是唯一正确恢复路径
    try:
        openid = scan_login.code2session(WECHAT_APPID, WECHAT_SECRET, code)
    except RuntimeError as e:
        msg = str(e)
        if "不可达" in msg or "返回异常" in msg:
            raise HTTPException(502, msg)
        raise HTTPException(401, msg)
    username = auth.find_user_by_openid(openid) \
        or auth.create_wechat_user(openid)
    if not scan_login.confirm_ticket(ticket, openid, username):
        raise HTTPException(409, "该二维码已使用，请在电脑上刷新后重新扫码")
    return {"ok": True, "username": username}


@app.post("/api/auth/wechat/cancel")
def api_wechat_cancel(request: Request, payload: dict = Body(...)):
    """小程序端"不是我操作的"：作废票据，电脑端随即收到 cancelled。"""
    _rate_limit(f"wxcancel:ip:{_client_ip(request)}", 60)
    ticket = str(payload.get("ticket") or "").strip()
    if not scan_login.TICKET_RE.match(ticket):
        raise HTTPException(400, "票据格式不正确")
    # 取消失败（不存在/已确认/已过期）不影响安全：票据到期后自然作废
    scan_login.cancel_ticket(ticket)
    return {"ok": True}


@app.get("/api/me")
def api_me(user: str = Depends(require_user)):
    return {"username": user}


@app.post("/api/recognize")
def api_recognize(request: Request, image: UploadFile = File(...),
                  tiles: str = Form("0"), name: str = Form(""),
                  crop: str = Form(""), user: str = Depends(require_user)):
    t0 = time.time()
    # 识别会真实调用付费视觉接口，按用户/按 IP 双重限流
    _rate_limit(f"recog:user:{user}", 60)
    _rate_limit(f"recog:ip:{_client_ip(request)}", 200)
    try:
        raw = image.file.read(MAX_UPLOAD_BYTES + 1)
    except Exception as e:
        logger.warning("读取上传内容失败 user=%s: %s", user, e)
        raise HTTPException(400, "无法读取上传的图片")
    if not raw:
        raise HTTPException(400, "上传内容为空")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, f"图片过大，请压缩到 "
                                 f"{MAX_UPLOAD_BYTES // 1024 // 1024}MB 以内")

    im = None
    try:
        im = Image.open(io.BytesIO(raw))
        # 解码前先看尺寸：阻断"解压炸弹"（体积很小但像素极大的 PNG）
        w0, h0 = im.size
        if w0 <= 0 or h0 <= 0 or w0 * h0 > MAX_IMAGE_PIXELS:
            raise ValueError(f"图片像素过多 {w0}x{h0}")
        img = ImageOps.exif_transpose(im).convert("RGB")
    except Exception as e:
        # 对外只给通用文案（原始异常可能含服务器路径），细节落日志
        logger.warning("无法读取图片 user=%s: %s", user, e)
        raise HTTPException(400, "无法读取图片，请上传常见的 JPEG/PNG/WebP 图片")
    finally:
        if im is not None:
            try:
                im.close()
            except Exception:
                pass
    if max(img.size) > MAX_EDGE:
        img.thumbnail((MAX_EDGE, MAX_EDGE), Image.LANCZOS)
    if crop.strip():
        img = img.crop(_parse_crop(crop, *img.size))

    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=88)
    saved = buf.getvalue()
    fname = _sanitize_name(_fix_name(name or image.filename or "image.jpg"))

    def stream():
        books_payload = None
        for chunk in _recognize_stream(img, tiles == "1"):
            yield chunk
            if chunk.startswith("event: books"):
                books_payload = json.loads(chunk.split("data: ", 1)[1])
            elif chunk.startswith("event: error"):
                return
        if books_payload is None:
            return
        rid, rdir = _new_result_dir()
        (rdir / "image.jpg").write_bytes(saved)
        result = {
            "id": rid,
            "name": fname,
            "user_id": user,
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "width": img.size[0],
            "height": img.size[1],
            "image_url": f"/results/{rid}/image.jpg",
            "shelf": books_payload.get("shelf", ""),
            "books": books_payload["books"],
            "tiles": books_payload["tiles"],
            "stats": books_payload.get("stats", {}),
        }
        _write_json_atomic(rdir / "result.json", result)
        st = result["stats"]
        logger.info("识别完成 rid=%s user=%s tiles=%d 耗时=%.1fs 书架=%s crop=%s "
                    "整图=%s 切块=%s 补漏=%s",
                    rid, user, len(result["tiles"]), time.time() - t0,
                    result["shelf"], crop.strip() or "-",
                    st.get("frame_count", "-"), st.get("detail_count", "-"),
                    st.get("filled", "-"))
        yield _sse("done", _public(result))

    return StreamingResponse(
        stream(), media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.get("/api/results")
def api_results(user: str = Depends(require_user)):
    # _my_records 会 glob 并解析磁盘上全部用户的 result.json，代价随全局记录数线性增长，
    # 与调用者无关；单用户高频刷就是廉价放大向量，故按用户限流。
    _rate_limit(f"results:list:user:{user}", 60)
    out = []
    for d in _my_records(user):
        books = d.get("books")
        out.append({"id": d.get("id", ""), "name": d.get("name", ""),
                    "count": len(books) if isinstance(books, list) else 0,
                    "shelf": d.get("shelf", ""),
                    "created_at": d.get("created_at", "")})
    return out


@app.get("/api/results/{rid}")
def api_result_detail(rid: str, user: str = Depends(require_user)):
    _rate_limit(f"results:detail:user:{user}", 60)
    d = _load_result(rid)
    if d is None or d.get("user_id") != user:
        raise HTTPException(404, "not found")
    return _public(d)


def _delete_result_dir(rid: str) -> bool:
    """删除整条记录目录。resolve 确认在 RESULTS_DIR 内，防符号链接/穿越。"""
    if not RID_RE.match(rid):
        return False
    root = RESULTS_DIR.resolve()
    rdir = (RESULTS_DIR / rid).resolve()
    if root not in rdir.parents or not rdir.is_dir():
        return False
    shutil.rmtree(rdir)
    return True


@app.delete("/api/results/{rid}")
def api_delete_result(rid: str, user: str = Depends(require_user)):
    """删除本人一条识别记录（不可逆）。"""
    _rate_limit(f"delete:user:{user}", 60)
    d = _load_result(rid)
    if d is None or d.get("user_id") != user:
        raise HTTPException(404, "not found")
    if not _delete_result_dir(rid):
        raise HTTPException(404, "not found")
    logger.info("删除记录 rid=%s user=%s", rid, user)
    return {"ok": True}


@app.post("/api/results/delete_all")
def api_delete_all_results(user: str = Depends(require_user)):
    """删除当前用户全部识别记录（不可逆操作，调用方需二次确认）。"""
    _rate_limit(f"deleteall:user:{user}", 10)
    n = 0
    for d in _my_records(user):
        if _delete_result_dir(d.get("id", "")):
            n += 1
    logger.info("删除全部记录 user=%s deleted=%d", user, n)
    return {"ok": True, "deleted": n}


@app.get("/results/{rid}/{fn}")
def serve_result_file(rid: str, fn: str, user: str = Depends(require_user)):
    # 图片是 <img>/小程序 image 直连的只读路由，同一照片页反复刷新/预览/重试属于正常行为，
    # 阈值放宽到 300/min（5/s），既能兜住脚本级刷取，又不会误拦正常浏览。
    _rate_limit(f"results:file:user:{user}", 300)
    if not RID_RE.match(rid) or fn != "image.jpg":
        raise HTTPException(404, "not found")
    d = _load_result(rid)
    if d is None or d.get("user_id") != user:
        raise HTTPException(404, "not found")
    p = (RESULTS_DIR / rid / fn).resolve()
    root = RESULTS_DIR.resolve()
    # 解析后必须仍在 results 目录内（防符号链接/穿越），且必须是普通文件
    if root not in p.parents or not p.is_file():
        raise HTTPException(404, "not found")
    return FileResponse(p, media_type="image/jpeg")


def _xl_safe(v):
    """Excel 公式注入防护：以 = + - @ 或制表符/回车开头的字符串前置单引号。"""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


@app.get("/api/export")
def api_export(ids: str = "", user: str = Depends(require_user)):
    """把指定识别记录导出为 Excel。ids 为逗号分隔的 rid，仅导出本人记录。"""
    _rate_limit(f"export:user:{user}", 60)
    # 先做参数长度校验：不限制的话一个超长 ids 会触发几十万次磁盘读取
    if len(ids) > 4000:
        raise HTTPException(400, "导出参数过长")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "书目"
    headers = ["图片名称", "书架号", "序号", "书名", "作者", "置信度",
               "x1", "y1", "x2", "y2"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="2E4B3F")
    n = 0
    for rid in [s.strip() for s in ids.split(",") if s.strip()][:MAX_EXPORT_IDS]:
        d = _load_result(rid)
        if d is None or d.get("user_id") != user:
            continue
        shelf = d.get("shelf", "")
        books = d.get("books")
        if not isinstance(books, list):
            continue
        for b in books:
            if not isinstance(b, dict):
                continue
            # bbox 可能缺失/不足 4 个/非数值：统一补齐到 4 列，保证列不错位
            bbox = [float(v) if isinstance(v, (int, float)) and not isinstance(v, bool)
                    else "" for v in (b.get("bbox") or [])[:4]]
            bbox += [""] * (4 - len(bbox))
            try:
                conf = float(b.get("confidence", 0))
            except (TypeError, ValueError):
                conf = 0.0
            ws.append([_xl_safe(str(d.get("name") or "")),
                       _xl_safe(str(b.get("shelf") or shelf or "")),
                       b.get("index", ""), _xl_safe(str(b.get("title") or "")),
                       _xl_safe(str(b.get("author") or "")),
                       conf, *bbox])
            n += 1
    if n == 0:
        raise HTTPException(404, "没有可导出的数据")
    for col, wdt in zip("ABCDEFGHIJ", (30, 10, 6, 40, 14, 8, 8, 8, 8, 8)):
        ws.column_dimensions[col].width = wdt
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument"
                   ".spreadsheetml.sheet",
        headers={"Content-Disposition":
                 "attachment; filename=books.xlsx; "
                 "filename*=UTF-8''%E4%B9%A6%E7%9B%AE.xlsx"},
    )


@app.post("/api/search")
def api_search(payload: dict = Body(...), user: str = Depends(require_user)):
    """多书搜索：在当前用户所有记录的书籍中做模糊匹配。"""
    queries = _check_search_queries(payload)
    _rate_limit(f"search:user:{user}", 60)
    return {"results": search_records(_my_records(user), queries)}


def _check_search_queries(payload: dict) -> list[str]:
    """公共参数校验：queries 需为 1~50 个合法书名，否则抛 400。"""
    queries = payload.get("queries")
    if not isinstance(queries, list) or not (1 <= len(queries) <= 50):
        raise HTTPException(400, "queries 需为 1~50 个书名")
    # 截断单个书名：search_records 里 difflib 是 O(len(q)*len(t))，
    # 不设上限的话一个 1MB 的 query 就能把 CPU 打满
    queries = [str(q).strip()[:MAX_QUERY_LEN] for q in queries]
    if any(not q for q in queries):
        raise HTTPException(400, "queries 不能包含空书名")
    return queries


@app.post("/api/search/export")
def api_search_export(payload: dict = Body(...), user: str = Depends(require_user)):
    """把多书搜索的"找到/未找到"汇总导出为 Excel。每行一个拟查找书名。"""
    queries = _check_search_queries(payload)
    _rate_limit(f"search:user:{user}", 60)
    results = search_records(_my_records(user), queries)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "查书结果"
    headers = ["拟查找书名", "是否找到", "推荐匹配书名", "所在记录", "书架号",
               "位置", "置信度", "匹配度"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="2E4B3F")
    for r in results:
        best = r.get("best")
        found = best is not None
        ws.append([
            _xl_safe(str(r.get("query") or "")),
            "找到" if found else "未找到",
            _xl_safe(str(best.get("title") or "")) if found else "",
            _xl_safe(str(best.get("record_name") or "")) if found else "",
            _xl_safe(str(best.get("shelf") or "")) if found else "",
            f"第 {best.get('index', '')} 本" if found else "",
            float(best.get("confidence") or 0) if found else "",
            float(best.get("score") or 0) if found else "",
        ])
    for col, wdt in zip("ABCDEFGH", (26, 10, 34, 22, 16, 10, 9, 9)):
        ws.column_dimensions[col].width = wdt
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument"
                   ".spreadsheetml.sheet",
        headers={"Content-Disposition":
                 "attachment; filename=books.xlsx; "
                 "filename*=UTF-8''%E6%9F%A5%E4%B9%A6%E7%BB%93%E6%9E%9C.xlsx"},
    )


@app.get("/api/providers")
def api_providers(user: str = Depends(require_user)):
    """列出可用的视觉模型供应商（供"重新识别"对话框勾选），不回传密钥。"""
    out = []
    for name in PROVIDER_DEFS:
        p = PROVIDERS[name]
        out.append({
            "name": name, "label": p["label"], "model": p["model"],
            "note": p["note"], "available": p["available"],
        })
    return {"providers": out, "default": DEFAULT_PROVIDER}


@app.post("/api/reidentify")
def api_reidentify(payload: dict = Body(...), user: str = Depends(require_user)):
    """对某条已有记录（该图）重新识别，支持选多个供应商并可开启切块、智能融合。

    body: {rid, providers: [name...], use_tiles: bool}
    结果会覆盖保存到该 rid，返回更新后的记录 + 各家识别本数 + 失败信息。
    """
    rid = str(payload.get("rid") or "").strip()
    use_tiles = payload.get("use_tiles") in (True, 1, "1", "true", "True")
    names = payload.get("providers")
    if not isinstance(names, list) or not names:
        names = [DEFAULT_PROVIDER]
    names = [n for n in (str(x).strip() for x in names) if n in PROVIDER_DEFS]
    if not names:
        raise HTTPException(400, "没有可用的识别供应商")
    _rate_limit(f"reidentify:user:{user}", 20)

    if not RID_RE.match(rid):
        raise HTTPException(404, "not found")
    rec = _load_result(rid)
    if rec is None or rec.get("user_id") != user:
        raise HTTPException(404, "not found")

    img_path = RESULTS_DIR / rid / "image.jpg"
    if not img_path.is_file():
        raise HTTPException(404, "没有找到该记录的原始图片")
    try:
        img = Image.open(io.BytesIO(img_path.read_bytes())).convert("RGB")
    except Exception as e:
        logger.warning("重识别无法读取图片 rid=%s: %s", rid, e)
        raise HTTPException(400, "无法读取该记录的原始图片")

    lists: list[dict] = []          # 各家组装后的书籍列表
    counts: dict[str, int] = {}     # name -> 识别本数
    failed: dict[str, str] = {}     # name -> 错误信息
    for name in names:
        if not PROVIDERS[name]["available"]:
            failed[name] = "未配置 API Key"
            continue
        try:
            detail, frame = _recognize_by_call(img, use_tiles, _provider_call(name))
            books, _shelf, _stats = _assemble(detail, frame)
            lists.append(books)
            counts[name] = len(books)
        except Exception as e:
            failed[name] = str(e)[:150]
            logger.warning("重识别供应商 %s 失败 rid=%s: %s", name, rid, e)

    if not lists:
        raise HTTPException(502, "所有选中的识别供应商均失败")
    # 单个供应商未选；先整体去重/排行/架标
    if len(lists) >= 2:
        try:
            books = fuse_book_lists(lists)
            books = dedup_row_x(books)
            books = snap_rows(books)
        except Exception as e:
            logger.warning("识别融合失败，退回首选结果: %s", e)
            books = list(lists[0])
    else:
        books = list(lists[0])
    assign_shelf_per_row(books)
    shelf_no = majority_shelf(books) or rec.get("shelf", "")
    for i, b in enumerate(books, 1):
        b["index"] = i

    rec["shelf"] = shelf_no
    rec["books"] = books
    rec["stats"] = {**rec.get("stats", {}),
                    "reidentified_by": names,
                    "reidentified_counts": counts,
                    "reidentified_failed": failed,
                    "reidentified_use_tiles": use_tiles}
    rec["reidentified_at"] = datetime.now().isoformat(timespec="seconds")
    _write_json_atomic(RESULTS_DIR / rid / "result.json", rec)
    logger.info("重识别完成 rid=%s user=%s 供应商=%s 融合=%d 本 失败=%s",
                rid, user, names, len(books), failed or "-")
    return {"result": _public(rec), "counts": counts, "failed": failed}


@app.get("/api/health")
def health():
    # 未鉴权端点：只回存活状态，不暴露模型名与密钥配置情况
    return {"ok": True}


# 生产模式：托管前端构建产物
DIST = BASE_DIR.parent / "web" / "dist"
if DIST.exists():
    from fastapi.staticfiles import StaticFiles
    app.mount("/", StaticFiles(directory=str(DIST), html=True), name="web")


if __name__ == "__main__":
    import uvicorn
    print(f"模型: {MODEL}  服务: http://127.0.0.1:8000")
    uvicorn.run(app, host="127.0.0.1", port=8000)
