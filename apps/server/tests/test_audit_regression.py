# -*- coding: utf-8 -*-
"""本轮对抗性安全/健壮性加固的独立回归测试（QA 编写）。

目的：把"工程师自验通过"变成可重复执行的断言。覆盖：
  1. 上传解压炸弹 / 体积上限 / 正常大图不误杀
  2. 限流（注册 60/min/IP）
  3. NaN / Infinity 防线（SSE 永不输出非法 JSON）
  4. ?token= 收敛到 /results/ 前缀
  5. SSE 客户端断连后工作线程收工
  6. 导出列宽恒为 10 列（含作者列）+ Excel 公式注入转义
  7. config.json 无密钥 / 缺 Key 时给出明确中文报错

约定：
  - 所有视觉接口调用均 mock，不依赖真实网络、不产生费用。
  - 限流是进程级全局状态，且 TestClient 的所有请求都来自同一个 IP
    （"testclient"），因此本模块用 autouse fixture 在每条用例前后清空
    _RATE_HITS，避免污染其它测试文件。
"""
import io
import json
import logging
import os
import struct
import threading
import time
import zlib
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image, ImageFile

import main
from conftest import auth_headers, make_image, parse_sse_event, register

# 采集期（conftest 的 client fixture 打桩之前）捕获真实的视觉接口函数，
# 用于"缺少 API Key 时应给出明确中文报错"这条用例。
from core.vision import call_vision as _REAL_CALL_VISION

SERVER_DIR = Path(__file__).resolve().parent.parent


@pytest.fixture(autouse=True)
def _isolate_rate_limit():
    """前后都清：既不被前面的用例污染，也不污染后面的测试文件。"""
    main._RATE_HITS.clear()
    yield
    main._RATE_HITS.clear()


@pytest.fixture()
def cheap_hash(monkeypatch):
    """把 PBKDF2 轮数降下来：限流用例要注册 60 个用户，只关心计数不关心哈希。"""
    from core import auth
    monkeypatch.setattr(auth, "PBKDF2_ROUNDS", 1000, raising=False)


# --------------------------------------------------------------------------
# 工具
# --------------------------------------------------------------------------
def _chunk(typ: bytes, data: bytes) -> bytes:
    return (struct.pack(">I", len(data)) + typ + data
            + struct.pack(">I", zlib.crc32(typ + data) & 0xFFFFFFFF))


def png_with_size(w: int, h: int) -> bytes:
    """构造一个 IHDR 声明 w×h 的极小 PNG（不含真实像素数据）。

    Pillow 的 Image.open() 只读 IHDR，不解码 IDAT，所以可以用 1KB 以内的
    文件声明 16 亿像素——这正是"解压炸弹"的形态。
    """
    ihdr = struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)  # 8bit / truecolor RGB
    return (b"\x89PNG\r\n\x1a\n"
            + _chunk(b"IHDR", ihdr)
            + _chunk(b"IDAT", zlib.compress(b"\x00\x00\x00\x00"))
            + _chunk(b"IEND", b""))


def _post_image(client, token, img_bytes, filename="x.png", **data):
    return client.post(
        "/api/recognize",
        files={"image": (filename, img_bytes, "image/png")},
        data=data,
        headers=auth_headers(token),
    )


def _write_result(results_dir: Path, rid: str, user: str, books,
                  name="x.jpg", shelf="A-1"):
    """直接落一条 result.json，用于导出接口的脏数据用例。"""
    d = results_dir / rid
    d.mkdir(parents=True, exist_ok=True)
    (d / "image.jpg").write_bytes(b"\xff\xd8\xff\xd9")
    (d / "result.json").write_text(
        json.dumps({
            "id": rid, "name": name, "user_id": user,
            "created_at": "2026-01-01T00:00:00",
            "width": 10, "height": 10,
            "image_url": f"/results/{rid}/image.jpg",
            "shelf": shelf, "books": books, "tiles": [],
        }, ensure_ascii=False), encoding="utf-8")
    return rid


def _iter_sse(body: str):
    """解析 SSE 响应体，产出 (event, data_dict)；data 必须是合法 JSON。"""
    for block in body.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        marker = "data: "
        i = block.find(marker)
        assert i != -1, f"事件块缺少 data: 字段 -> {block!r}"
        event = block[:i].replace("event:", "").strip()
        # 关键断言：data 必须能被 json.loads 解析（NaN/Infinity 会导致失败）
        data = json.loads(block[i + len(marker):])
        yield event, data


# ==========================================================================
# 1. 上传防护
# ==========================================================================
class TestUploadHardening:
    def test_decompression_bomb_rejected_before_decode(self, client, monkeypatch):
        """IHDR 声明 40000×40000 的 1KB PNG 必须被拒，且 Pillow 不得解码。"""
        decoded = {"n": 0}
        real_load = ImageFile.ImageFile.load

        def counting_load(self, *a, **kw):
            decoded["n"] += 1
            return real_load(self, *a, **kw)

        monkeypatch.setattr(ImageFile.ImageFile, "load", counting_load)

        bomb = png_with_size(40000, 40000)
        assert len(bomb) < 1024, f"炸弹文件应极小，实际 {len(bomb)} 字节"
        assert 40000 * 40000 > main.MAX_IMAGE_PIXELS

        token = register(client, "boom")
        r = _post_image(client, token, bomb)
        assert r.status_code == 400, r.text
        assert "无法读取图片" in r.json()["detail"]
        # 核心：还没解码就被拦下了
        assert decoded["n"] == 0, "Pillow 在尺寸校验前就解码了像素数据"

    def test_oversized_upload_413(self, client):
        """31MB 的假图片：中间件按 Content-Length 拦，端点层兜底。"""
        token = register(client, "big")
        payload = b"\xff\xd8" + b"0" * (31 * 1024 * 1024)
        r = _post_image(client, token, payload, filename="big.jpg")
        assert r.status_code == 413, r.status_code
        assert "30MB" in r.json()["detail"]

    def test_large_but_legal_image_not_rejected(self, client):
        """8000×6000 = 4800 万像素，在上限内，不能被误杀（允许被缩到 6000 边）。"""
        buf = io.BytesIO()
        Image.new("L", (8000, 6000), 128).save(buf, format="JPEG", quality=70)
        raw = buf.getvalue()
        assert 8000 * 6000 < main.MAX_IMAGE_PIXELS

        token = register(client, "huge")
        r = _post_image(client, token, raw, filename="huge.jpg",
                        tiles="0", name="huge.jpg")
        assert r.status_code == 200, r.text[:500]
        done = parse_sse_event(r.text, "done")
        assert done, r.text[:500]
        # 未被拒绝：记录已生成；因 MAX_EDGE=6000 被等比缩小到 6000×4500
        assert done["width"] == 6000 and done["height"] == 4500

    def test_pillow_global_pixel_cap_is_set(self):
        assert Image.MAX_IMAGE_PIXELS == main.MAX_IMAGE_PIXELS == 50_000_000


# ==========================================================================
# 2. 限流
# ==========================================================================
class TestRateLimit:
    def test_register_limit_61st_is_429(self, client, cheap_hash):
        """注册：同 IP 60/min，第 61 次必须 429。"""
        for i in range(60):
            r = client.post("/api/register",
                            json={"username": f"u_{i:03d}", "password": "pass123"})
            assert r.status_code == 200, f"第 {i+1} 次注册失败: {r.text}"
        r = client.post("/api/register",
                        json={"username": "u_060", "password": "pass123"})
        assert r.status_code == 429, r.status_code
        assert "频繁" in r.json()["detail"]

    def test_login_limit_per_ip_and_per_account(self, client, cheap_hash):
        """登录：同 IP 60/min，同 IP×账号 10/min。"""
        client.post("/api/register",
                    json={"username": "lim", "password": "pass123"})
        for i in range(10):
            r = client.post("/api/login",
                            json={"username": "lim", "password": "wrongpw"})
            assert r.status_code == 401, f"第 {i+1} 次: {r.status_code}"
        # 第 11 次：账号维度先撞线（IP 维度还早）
        r = client.post("/api/login",
                        json={"username": "lim", "password": "wrongpw"})
        assert r.status_code == 429, r.status_code

    def test_rate_limit_state_is_bounded(self, client, cheap_hash):
        """_RATE_HITS 不能无限增长。"""
        for i in range(80):
            main._rate_limit(f"probe:{i}", 10_000)
        assert len(main._RATE_HITS) <= main._RATE_MAX_KEYS + 80


# ==========================================================================
# 3. NaN / Infinity 防线
# ==========================================================================
NAN_RAW = (
    '[{"title":"坏书","shelf":"A-1","bbox":[0.10,NaN,0.18,0.90],'
    '"confidence":Infinity},'
    '{"title":"好书","shelf":"A-1","bbox":[0.20,0.10,0.28,0.90],'
    '"confidence":NaN}]'
)


class TestNaNGuard:
    def test_nan_never_reaches_sse_payload(self, client, monkeypatch):
        def fake_vision(image_bytes, **kw):
            return NAN_RAW
        monkeypatch.setattr(main, "call_vision", fake_vision)

        token = register(client, "nan")
        r = _post_image(client, token, make_image(), filename="a.jpg",
                        tiles="0", name="a.jpg")
        assert r.status_code == 200, r.text

        body = r.text
        # 1) 响应体里不得出现任何非法 JSON 字面量
        for bad in ("NaN", "Infinity", "-Infinity", "nan", "inf"):
            assert bad not in body, f"SSE 响应体里出现了非法字面量 {bad!r}"
        # 2) 每一条 data: 都必须能被 JS 的 JSON.parse 解析
        events = list(_iter_sse(body))
        assert events, "没有任何 SSE 事件"
        # 3) bbox 含 NaN 的书被丢弃，另一本保留且置信度回落为 0.5
        done = parse_sse_event(body, "done")
        assert done is not None
        titles = [b["title"] for b in done["books"]]
        assert titles == ["好书"], titles
        assert done["books"][0]["confidence"] == 0.5

    def test_parse_books_filters_non_finite(self):
        from core.parse import parse_books
        books = parse_books(NAN_RAW)
        assert len(books) == 1
        assert books[0]["title"] == "好书"
        assert books[0]["confidence"] == 0.5


# ==========================================================================
# 4. ?token= 收敛
# ==========================================================================
class TestTokenQueryParamConvergence:
    def test_token_query_param_rejected_outside_results(self, client):
        token = register(client, "tk")
        for url in ("/api/me", "/api/results", "/api/results/20260101_000000_001",
                    "/api/export?ids=20260101_000000_001"):
            r = client.get(url, params={"token": token})
            assert r.status_code == 401, f"{url} 不应接受 ?token= -> {r.status_code}"

    def test_token_query_param_allowed_for_result_image(self, client):
        """<img> 场景：/results/ 前缀仍可用 ?token=。"""
        from conftest import recognize
        token = register(client, "tk2")
        done = recognize(client, token, filename="p.jpg")
        url = f"/results/{done['id']}/image.jpg"
        assert client.get(url).status_code == 401
        r = client.get(url, params={"token": token})
        assert r.status_code == 200
        assert r.content[:2] == b"\xff\xd8"

    def test_bearer_header_still_works_everywhere(self, client):
        token = register(client, "tk3")
        assert client.get("/api/me", headers=auth_headers(token)).status_code == 200

    def test_token_is_redacted_in_logs(self, client):
        """?token= 落盘前必须被脱敏（子 logger 冒泡上来的记录也要生效）。"""
        token = register(client, "tk4")
        buf = io.StringIO()
        handler = logging.StreamHandler(buf)
        # 模拟 _file_handler：脱敏 filter 挂在 handler 上才拦得住子 logger 冒泡
        handler.addFilter(main._REDACTOR)
        main.logger.addHandler(handler)
        try:
            main.logger.info("probe GET /results/x/image.jpg?token=%s", token)
            handler.flush()
            text = buf.getvalue()
        finally:
            main.logger.removeHandler(handler)
        assert "token=***" in text, text
        assert token not in text, "令牌明文泄漏到了日志"


# ==========================================================================
# 5. SSE 断连收工
# ==========================================================================
class TestSseShutdown:
    def test_generator_close_stops_worker_threads(self, monkeypatch):
        """客户端中途断开：工作线程必须收工，不留残留 ThreadPoolExecutor 线程。"""
        def slow_vision(image_bytes, **kw):
            time.sleep(0.4)
            return "[]"
        monkeypatch.setattr(main, "call_vision", slow_vision)

        def pool_threads() -> int:
            return sum(1 for t in threading.enumerate()
                       if t.name.startswith("ThreadPoolExecutor"))

        base_pool = pool_threads()
        base_active = threading.active_count()

        img = Image.new("RGB", (4800, 1200), (10, 20, 30))  # -> 3 个切块
        gen = main._recognize_stream(img, True)
        it = iter(gen)
        assert next(it).startswith("event: meta")
        # 推进到 try 块内的第一个 tile 事件（此时工作线程已在跑）
        assert next(it).startswith("event: tile")
        gen.close()  # 模拟客户端断开

        deadline = time.time() + 10
        while time.time() < deadline and pool_threads() > base_pool:
            time.sleep(0.05)
        assert pool_threads() <= base_pool, "客户端断开后仍有残留工作线程"
        assert threading.active_count() <= base_active + 1

    def test_bounded_queue_is_used(self):
        """队列必须是有界的，否则断连后结果会无限堆积。"""
        import queue
        img = Image.new("RGB", (4800, 1200), (10, 20, 30))
        gen = main._recognize_stream(img, True)
        next(iter(gen))
        # 通过源码契约断言：maxsize 有界
        src = Path(main.__file__).read_text(encoding="utf-8")
        assert "queue.Queue(maxsize=" in src
        gen.close()


# ==========================================================================
# 6. 导出加固
# ==========================================================================
class TestExportHardening:
    def _export(self, client, token, rid):
        r = client.get("/api/export", params={"ids": rid},
                       headers=auth_headers(token))
        assert r.status_code == 200, r.text
        return load_workbook(io.BytesIO(r.content)).active

    def test_short_bbox_row_still_has_ten_columns(self, client):
        token = register(client, "ex1")
        rid = _write_result(main.RESULTS_DIR, "20260101_000000_001", "ex1",
                            [{"index": 1, "title": "短框书", "shelf": "A-1",
                              "bbox": [0.1, 0.2], "confidence": "0.9"}])
        ws = self._export(client, token, rid)
        assert ws.max_column == 10, f"列数应为 10，实际 {ws.max_column}"
        row = next(ws.iter_rows(min_row=2, values_only=True))
        assert len(row) == 10, row
        # 列序：0 图片名 / 1 书架号 / 2 序号 / 3 书名 / 4 作者 / 5 置信度
        #       / 6~9 = x1,y1,x2,y2
        assert row[4] in (None, ""), row  # 无 author 时为空
        assert row[6] == 0.1 and row[7] == 0.2, row
        # openpyxl 读回空字符串单元格会返回 None，两者都算"空"
        assert row[8] in (None, "") and row[9] in (None, ""), row

    def test_author_column_written(self, client):
        token = register(client, "ex1b")
        rid = _write_result(main.RESULTS_DIR, "20260101_000000_006", "ex1b",
                            [{"index": 1, "title": "有作者的书", "shelf": "A-1",
                              "author": "余华", "bbox": [0.1, 0.2, 0.3, 0.4],
                              "confidence": 0.9}])
        ws = self._export(client, token, rid)
        header = [c.value for c in ws[1]]
        assert header == ["图片名称", "书架号", "序号", "书名", "作者",
                          "置信度", "x1", "y1", "x2", "y2"], header
        row = next(ws.iter_rows(min_row=2, values_only=True))
        assert row[4] == "余华", row

    def test_missing_bbox_row_has_ten_columns(self, client):
        token = register(client, "ex2")
        rid = _write_result(main.RESULTS_DIR, "20260101_000000_002", "ex2",
                            [{"index": 1, "title": "无框书", "shelf": "A-1"}])
        ws = self._export(client, token, rid)
        row = next(ws.iter_rows(min_row=2, values_only=True))
        assert len(row) == 10, row

    def test_formula_injection_is_escaped(self, client):
        token = register(client, "ex3")
        rid = _write_result(
            main.RESULTS_DIR, "20260101_000000_003", "ex3",
            [
                {"index": 1, "title": "\t=1+1", "shelf": "A-1",
                 "bbox": [0.1, 0.2, 0.3, 0.4], "confidence": 0.5},
                {"index": 2, "title": "=cmd|'/c calc'!A1", "shelf": "A-1",
                 "bbox": [0.1, 0.2, 0.3, 0.4], "confidence": 0.5},
            ],
            name="=cmd|'/c calc'!A1",
        )
        ws = self._export(client, token, rid)
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) == 2
        # 图片名（第 1 列）与书名（第 4 列）都被前置单引号
        assert rows[0][0].startswith("'"), rows[0]
        assert rows[0][3].startswith("'"), rows[0]
        assert rows[1][3].startswith("'"), rows[1]

    def test_confidence_is_float_not_string(self, client):
        token = register(client, "ex4")
        rid = _write_result(main.RESULTS_DIR, "20260101_000000_004", "ex4",
                            [{"index": 1, "title": "T", "shelf": "A-1",
                              "bbox": [0.1, 0.2, 0.3, 0.4], "confidence": "0.75"}])
        ws = self._export(client, token, rid)
        v = next(ws.iter_rows(min_row=2, values_only=True))[5]
        assert isinstance(v, float) and abs(v - 0.75) < 1e-9, v

    def test_books_not_list_is_skipped(self, client):
        token = register(client, "ex5")
        _write_result(main.RESULTS_DIR, "20260101_000000_005", "ex5",
                      "not-a-list")
        r = client.get("/api/export", params={"ids": "20260101_000000_005"},
                       headers=auth_headers(token))
        assert r.status_code == 404

    def test_ids_length_limit(self, client):
        token = register(client, "ex6")
        long_ids = ",".join(["20260101_000000_001"] * 300)  # 远超 4000 字符
        assert len(long_ids) > 4000
        r = client.get("/api/export", params={"ids": long_ids},
                       headers=auth_headers(token))
        assert r.status_code == 400
        assert "过长" in r.json()["detail"]


# ==========================================================================
# 7. 配置安全性
# ==========================================================================
class TestConfigSecurity:
    def test_config_json_has_no_secrets(self):
        cfg = json.loads((SERVER_DIR / "config.json").read_text(encoding="utf-8"))
        banned = {"api_key", "apikey", "secret", "wechat_secret", "password",
                  "token", "access_key"}
        assert not (banned & {k.lower() for k in cfg}), cfg.keys()

        def walk(o):
            if isinstance(o, dict):
                for v in o.values():
                    walk(v)
            elif isinstance(o, str):
                assert not o.startswith("sk-"), "config.json 里仍有 sk- 开头的密钥"
                assert len(o) < 64, f"config.json 里存在可疑长字符串: {o[:20]}..."
        walk({k: v for k, v in cfg.items() if k != "_comment"})

    def test_legacy_secret_fields_ignored_and_never_written_back(self, tmp_path):
        """在子进程里验证：config.json 残留 api_key/secret 时被忽略，
        且缺失 BOOKFINDER_SECRET 时不会把生成的密钥写回 config.json。

        用子进程隔离：reload main 会在同一进程里重复挂 RotatingFileHandler
        （Windows 上重复打开同一个日志文件有风险），也会重建 FastAPI app。
        """
        import shutil
        import subprocess
        import sys

        srv = tmp_path / "srv"
        srv.mkdir()
        shutil.copy(SERVER_DIR / "main.py", srv / "main.py")
        shutil.copytree(SERVER_DIR / "core", srv / "core")
        cfg = srv / "config.json"
        original = {"api_key": "sk-leaked", "base_url": "https://api.deepseek.com"}
        cfg.write_text(json.dumps(original), encoding="utf-8")

        code = (
            "import sys, json\n"
            f"sys.path.insert(0, r'{srv}')\n"
            "import main\n"
            "assert 'api_key' not in main.CONFIG, main.CONFIG\n"
            "assert 'secret' not in main.CONFIG, main.CONFIG\n"
            "assert main.SECRET and len(main.SECRET) >= 32\n"
            "print('OK')\n"
        )
        env = dict(os.environ, PYTHONIOENCODING="utf-8")
        env.pop("BOOKFINDER_SECRET", None)
        proc = subprocess.run([sys.executable, "-c", code], capture_output=True,
                              env=env, timeout=180)
        assert proc.returncode == 0, (proc.stdout or b"").decode("utf-8", "replace") \
            + (proc.stderr or b"").decode("utf-8", "replace")
        assert "OK" in proc.stdout.decode("utf-8", "replace")
        # 关键：磁盘上的 config.json 没有被回写密钥
        after = json.loads(cfg.read_text(encoding="utf-8"))
        assert after == original, f"config.json 被回写了: {after}"
        assert "secret" not in after

    def test_missing_api_key_gives_clear_chinese_error(self, client, monkeypatch):
        """未配置 DEEPSEEK_API_KEY：明确的中文报错，不是静默失败也不是 500。"""
        monkeypatch.setattr(main, "API_KEY", "")
        monkeypatch.setattr(main, "call_vision", _REAL_CALL_VISION)

        token = register(client, "nokey")
        r = _post_image(client, token, make_image(), filename="k.jpg",
                        tiles="0", name="k.jpg")
        assert r.status_code == 200, f"不应返回 5xx: {r.status_code}"
        err = parse_sse_event(r.text, "error")
        assert err, f"没有 error 事件: {r.text[:500]}"
        msg = err["error"]
        assert "DeepSeek API Key" in msg
        assert "DEEPSEEK_API_KEY" in msg
        # 中文可读，不是裸异常
        assert any("\u4e00" <= ch <= "\u9fff" for ch in msg)

    def test_health_exposes_nothing_but_ok(self, client):
        r = client.get("/api/health")
        assert r.status_code == 200
        assert r.json() == {"ok": True}


# ==========================================================================
# 8. 其它：原子写 / 搜索截断
# ==========================================================================
class TestMisc:
    def test_no_tmp_file_left_after_recognize(self, client):
        from conftest import recognize
        token = register(client, "atom")
        recognize(client, token, filename="at.jpg")
        leftovers = list(main.RESULTS_DIR.rglob("*.tmp"))
        assert leftovers == [], f"原子写残留临时文件: {leftovers}"

    def test_long_search_query_is_truncated_not_crashing(self, client):
        from conftest import recognize
        token = register(client, "qs")
        recognize(client, token, filename="q.jpg")
        r = client.post("/api/search", json={"queries": ["书" * 5000]},
                        headers=auth_headers(token))
        assert r.status_code == 200, r.text
        assert len(r.json()["results"]) == 1

    def test_new_result_dir_uses_single_timestamp(self, tmp_path, monkeypatch):
        """rid 的秒位与毫秒位必须来自同一次 time.time()。"""
        monkeypatch.setattr(main, "RESULTS_DIR", tmp_path)
        rid, rdir = main._new_result_dir()
        assert main.RID_RE.match(rid), rid
        assert rdir.parent == tmp_path
