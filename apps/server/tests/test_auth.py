# -*- coding: utf-8 -*-
"""用户认证与记录隔离测试。"""
import io

from openpyxl import load_workbook

from conftest import auth_headers, recognize, register


class TestAuth:
    def test_register_and_login(self, client):
        token = register(client, "alice")
        assert token
        r = client.post("/api/login",
                        json={"username": "alice", "password": "pass123"})
        assert r.status_code == 200
        assert r.json()["username"] == "alice"
        assert r.json()["token"]

    def test_duplicate_register(self, client):
        register(client, "alice")
        r = client.post("/api/register",
                        json={"username": "alice", "password": "other123"})
        assert r.status_code == 409
        assert r.json()["detail"] == "用户名已存在"

    def test_wrong_password(self, client):
        register(client, "alice")
        r = client.post("/api/login",
                        json={"username": "alice", "password": "wrong123"})
        assert r.status_code == 401
        assert r.json()["detail"] == "用户名或密码错误"

    def test_login_unknown_user(self, client):
        r = client.post("/api/login",
                        json={"username": "nobody", "password": "pass123"})
        assert r.status_code == 401

    def test_invalid_register_fields(self, client):
        r = client.post("/api/register",
                        json={"username": "a", "password": "pass123"})
        assert r.status_code == 400
        r = client.post("/api/register",
                        json={"username": "alice", "password": "123"})
        assert r.status_code == 400

    def test_me(self, client):
        token = register(client, "alice")
        r = client.get("/api/me", headers=auth_headers(token))
        assert r.status_code == 200
        assert r.json() == {"username": "alice"}

    def test_me_unauthorized(self, client):
        assert client.get("/api/me").status_code == 401
        assert client.get("/api/results").status_code == 401
        assert client.post("/api/search", json={"queries": ["x"]}).status_code == 401

    def test_bad_token(self, client):
        r = client.get("/api/me",
                       headers={"Authorization": "Bearer not-a-token"})
        assert r.status_code == 401

    def test_image_token_query_param(self, client):
        """<img> 场景：用 ?token= 访问记录图片。"""
        token = register(client, "alice")
        done = recognize(client, token)
        url = f"/results/{done['id']}/image.jpg"
        assert client.get(url).status_code == 401
        r = client.get(url, params={"token": token})
        assert r.status_code == 200
        assert r.content[:2] == b"\xff\xd8"  # JPEG


class TestIsolation:
    def test_results_isolated(self, client):
        ta = register(client, "alice")
        tb = register(client, "bob")
        done_a = recognize(client, ta, filename="a.jpg")
        done_b = recognize(client, tb, filename="b.jpg")

        ids_a = [d["id"] for d in client.get(
            "/api/results", headers=auth_headers(ta)).json()]
        ids_b = [d["id"] for d in client.get(
            "/api/results", headers=auth_headers(tb)).json()]
        assert ids_a == [done_a["id"]]
        assert ids_b == [done_b["id"]]

    def test_detail_cross_user_404(self, client):
        ta = register(client, "alice")
        tb = register(client, "bob")
        done_b = recognize(client, tb)
        r = client.get(f"/api/results/{done_b['id']}",
                       headers=auth_headers(ta))
        assert r.status_code == 404

    def test_image_cross_user_404(self, client):
        ta = register(client, "alice")
        tb = register(client, "bob")
        done_b = recognize(client, tb)
        r = client.get(f"/results/{done_b['id']}/image.jpg",
                       params={"token": ta})
        assert r.status_code == 404

    def test_export_only_own(self, client):
        ta = register(client, "alice")
        tb = register(client, "bob")
        done_a = recognize(client, ta, filename="alice的书架.jpg")
        done_b = recognize(client, tb, filename="bob的书架.jpg")

        # 导出他人记录：无数据
        r = client.get("/api/export", params={"ids": done_b["id"]},
                       headers=auth_headers(ta))
        assert r.status_code == 404

        # 混合 ids：只导出本人记录
        r = client.get("/api/export",
                       params={"ids": f"{done_a['id']},{done_b['id']}"},
                       headers=auth_headers(ta))
        assert r.status_code == 200
        ws = load_workbook(io.BytesIO(r.content)).active
        names = {row[0] for row in ws.iter_rows(min_row=2, values_only=True)}
        assert names == {"alice的书架.jpg"}

    def test_result_has_created_at(self, client):
        ta = register(client, "alice")
        done = recognize(client, ta)
        assert done["created_at"]
        assert "user_id" not in done
        d = client.get(f"/api/results/{done['id']}",
                       headers=auth_headers(ta)).json()
        assert d["created_at"] == done["created_at"]
        item = client.get("/api/results", headers=auth_headers(ta)).json()[0]
        assert item["created_at"] and item["shelf"] == "A-03"
