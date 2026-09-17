# -*- coding: utf-8 -*-
"""两遍识别（整图框架 + 切块细节）的流程测试：mock call_vision 按图片尺寸区分
框架块（整图）与细节块（切块），验证补漏、stats、pass 标记。"""
import io
import json

from PIL import Image

from conftest import auth_headers, register

DETAIL_RAW = json.dumps([
    {"title": "细节书", "author": "张三", "shelf": "A-9",
     "bbox": [0.02, 0.10, 0.10, 0.90], "confidence": 0.9},
], ensure_ascii=False)

FRAME_RAW = json.dumps([
    {"title": "细节书", "author": "张三", "shelf": "A-9",
     "bbox": [0.01, 0.10, 0.06, 0.90], "confidence": 0.9},
    {"title": "框架漏网书", "author": "李四", "shelf": "A-9",
     "bbox": [0.80, 0.10, 0.88, 0.90], "confidence": 0.9},
], ensure_ascii=False)


def _sse_events(text: str):
    """解析全部 SSE 事件为 [(event, data)]。"""
    events = []
    for block in text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        i = block.find("data: ")
        events.append((block[:i].replace("event:", "").strip(),
                       json.loads(block[i + 6:])))
    return events


def _make_big_image(w=2000, h=1200) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), (60, 90, 120)).save(buf, format="JPEG")
    return buf.getvalue()


class TestTwoPass:
    def _run(self, client, monkeypatch):
        def fake_vision(image_bytes, **kw):
            im = Image.open(io.BytesIO(image_bytes))
            return FRAME_RAW if im.size == (2000, 1200) else DETAIL_RAW
        monkeypatch.setattr("main.call_vision", fake_vision)

        token = register(client, "twopass")
        r = client.post("/api/recognize",
                        files={"image": ("big.jpg", _make_big_image(),
                                          "image/jpeg")},
                        data={"tiles": "1"}, headers=auth_headers(token))
        assert r.status_code == 200, r.text
        return _sse_events(r.text)

    def test_frame_pass_fills_missing_book(self, client, monkeypatch):
        events = self._run(client, monkeypatch)
        done = dict((e, d) for e, d in events)["done"]

        # 细节块 2 本 + 框架补漏 1 本 = 3 本
        assert len(done["books"]) == 3
        filled = [b for b in done["books"] if b.get("source") == "frame"]
        assert len(filled) == 1
        assert filled[0]["title"] == "框架漏网书"
        assert abs(filled[0]["confidence"] - 0.9 * 0.8) < 1e-6

        # stats：整图 2 本、切块 2 本、补漏 1 本
        stats = done["stats"]
        assert stats["frame_count"] == 2
        assert stats["detail_count"] == 2
        assert stats["filled"] == 1

    def test_frame_tile_marked_in_events(self, client, monkeypatch):
        events = self._run(client, monkeypatch)
        tiles = [d for e, d in events if e == "tile"]
        # 2000x1200 -> 2 个细节块 + 1 个整图框架块
        assert sum(1 for t in tiles if t["pass"] == "detail") == 2
        assert sum(1 for t in tiles if t["pass"] == "frame") == 1
        meta = dict((e, d) for e, d in events)["meta"]
        assert meta["tiles_total"] == 3

    def test_frame_same_title_displaced_not_filled(self, client, monkeypatch):
        """整图遍 bbox 定位不准（实测会把一排书整体右移）：与切块书同名
        （归一化相等/互为子串）且 y 带重叠的整图书是错位重复，不应补进结果。"""
        detail = json.dumps([
            {"title": "细节书", "author": "张三", "shelf": "A-9",
             "bbox": [0.02, 0.10, 0.10, 0.90], "confidence": 0.9},
        ], ensure_ascii=False)
        frame = json.dumps([
            # 与细节书同名但 bbox 错位（中心点规则兜不住）
            {"title": "细节书", "author": "张三", "shelf": "A-9",
             "bbox": [0.55, 0.10, 0.60, 0.90], "confidence": 0.9},
            {"title": "框架漏网书", "author": "李四", "shelf": "A-9",
             "bbox": [0.80, 0.10, 0.88, 0.90], "confidence": 0.9},
        ], ensure_ascii=False)

        def fake_vision(image_bytes, **kw):
            im = Image.open(io.BytesIO(image_bytes))
            return frame if im.size == (2000, 1200) else detail
        monkeypatch.setattr("main.call_vision", fake_vision)

        token = register(client, "twopass2")
        r = client.post("/api/recognize",
                        files={"image": ("big.jpg", _make_big_image(),
                                          "image/jpeg")},
                        data={"tiles": "1"}, headers=auth_headers(token))
        assert r.status_code == 200, r.text
        done = dict((e, d) for e, d in _sse_events(r.text))["done"]
        filled = [b for b in done["books"] if b.get("source") == "frame"]
        assert [b["title"] for b in filled] == ["框架漏网书"]
        assert done["stats"]["filled"] == 1

    def test_author_flows_to_result_and_export(self, client, monkeypatch):
        events = self._run(client, monkeypatch)
        done = dict((e, d) for e, d in events)["done"]
        authors = {b["title"]: b.get("author", "") for b in done["books"]}
        assert authors.get("细节书") == "张三"
        assert authors.get("框架漏网书") == "李四"

        from openpyxl import load_workbook
        token = client.post("/api/login", json={
            "username": "twopass", "password": "pass123"}).json()["token"]
        r = client.get("/api/export", params={"ids": done["id"]},
                       headers=auth_headers(token))
        assert r.status_code == 200
        ws = load_workbook(io.BytesIO(r.content)).active
        assert [c.value for c in ws[1]][4] == "作者"
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert any(row[4] == "李四" for row in rows)

    def test_tiles0_no_frame_pass(self, client, monkeypatch):
        """tiles=0 保持现状：只整图一遍，不重复调用。"""
        calls = []

        def fake_vision(image_bytes, **kw):
            calls.append(1)
            return DETAIL_RAW
        monkeypatch.setattr("main.call_vision", fake_vision)

        token = register(client, "single")
        r = client.post("/api/recognize",
                        files={"image": ("s.jpg", _make_big_image(),
                                          "image/jpeg")},
                        data={"tiles": "0"}, headers=auth_headers(token))
        assert r.status_code == 200, r.text
        assert len(calls) == 1
        events = _sse_events(r.text)
        assert all(d.get("pass") == "detail" for e, d in events if e == "tile")

    def test_small_image_no_frame_pass(self, client, monkeypatch):
        """小图只有 1 个切块时，框架块与细节块相同，不重复调用。"""
        calls = []

        def fake_vision(image_bytes, **kw):
            calls.append(1)
            return DETAIL_RAW
        monkeypatch.setattr("main.call_vision", fake_vision)

        token = register(client, "small")
        from conftest import make_image
        r = client.post("/api/recognize",
                        files={"image": ("s.jpg", make_image(), "image/jpeg")},
                        data={"tiles": "1"}, headers=auth_headers(token))
        assert r.status_code == 200, r.text
        assert len(calls) == 1
