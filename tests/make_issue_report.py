# -*- coding: utf-8 -*-
"""生成 V3 对抗性检查问题清单 Excel。"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# (级别, 端, 位置, 问题, 触发场景, 建议修法, 状态)
ISSUES = [
    # ---------- P0 ----------
    ("P0", "小程序+后端", "utils/request.js uploadRecognize / main.py stream()",
     "wx.uploadFile 默认 60s 超时，识别常需 1~3 分钟；超时后客户端报失败，且服务端在流末尾才落盘，已付费的识别结果直接丢弃",
     "真机识别任意需切块的照片，弱网下几乎必现",
     "改为任务制：POST 立即返回 rid，服务端先落盘后推送，小程序轮询 /api/results/{rid} 取结果",
     "未修复（需架构调整，待确认）"),
    # ---------- P1 已修复 ----------
    ("P1", "后端", "main.py _recognize_stream",
     "work() 的 try 只包了 call_vision，img.crop/save 异常会使 events 结束哨兵丢失，SSE 生成器永久阻塞，约 40 个并发坏请求拖死服务",
     "OOM 或 PIL 内部错误时",
     "run_all 外裹 try/finally 保证哨兵到达，work() try 扩到整个函数体", "已修复"),
    ("P1", "后端", "main.py _parse_crop",
     "合法 crop（面积≥1%）经像素取整后可退化为零宽图，img.save 抛 ValueError 裸 500",
     "小图 + 极端细长 crop（实测 3×100 图 crop 0.33~0.34 可复现）",
     "取整后校验像素 x1>x0 且 y1>y0，否则 400", "已修复"),
    ("P1", "后端", "core/auth.py verify_token",
     "非 ASCII 字符的畸形 token 使 payload.encode('ascii') 抛 UnicodeEncodeError → 500",
     "攻击者发送含非 ASCII 字节的 Authorization 头",
     "encode 包 try，异常返回 None 按无效 token 处理", "已修复"),
    ("P1", "后端", "main.py config 写回",
     "secret 首次生成后直接 write_text 覆写 config.json，写盘中途崩溃会截断文件、连同 api_key 一起丢失",
     "首次启动时进程被杀/断电",
     "临时文件 + os.replace 原子替换", "已修复"),
    ("P1", "后端", "core/search.py best 语义",
     "契约写 best=置信度最高，代码取的是 score(0.6相似度+0.4置信度) 最高，与用户'推荐置信度最高的那一本'的要求不符",
     "相似度满分但低置信的书压过高置信的正确书",
     "best = max(matches, key=(confidence, score))，matches 仍按 score 排序，文档/测试同步", "已修复"),
    ("P1", "小程序", "pages/index/index.js onConfirmCrop",
     "双重裁剪：canvas 已裁出局部图，又把相对原图的 crop 传给服务端再裁一次，识别区域系统性错误（架标、目标书被丢掉）",
     "canvas 裁剪成功时（开发者工具中几乎总是成功）",
     "裁剪成功传 crop=''；仅回退原图时传归一化 crop", "已修复"),
    ("P1", "小程序", "utils/request.js",
     "所有 401 一律清缓存+reLaunch 登录页+提示'登录已过期'，密码输错（也是 401）会重载页面、清空输入、提示错误",
     "登录页输错密码",
     "request/uploadFile 增加 skipAuth 选项，login/register/wechat 三处跳过 401 拦截", "已修复"),
    ("P1", "Web", "components/CropEditor.tsx",
     "图片未加载时 overlay 尺寸为 0，拖动产生 NaN 坐标且状态无法自愈，confirm 会把 'NaN,NaN,NaN,NaN' 发给服务端",
     "慢网速下图片加载中拖动裁剪框（角把手在 0 尺寸框外仍点得到）",
     "onPointerMove 加 0 尺寸守卫；confirm 前 Number.isFinite 校验，异常按整图处理", "已修复"),
    ("P1", "Web", "components/UploadZone.tsx",
     "识别进度按文件名匹配，同批两个同名文件（或一秒内连拍两张同名照片）进度互相串扰、状态错乱",
     "同名文件批量上传",
     "每个任务 crypto.randomUUID() 唯一 id，进度与 key 都用 id", "已修复"),
    ("P1", "Web", "components/SearchPanel.tsx",
     "zod 校验是死代码（表单无 submit），超长输入不经阻断直接进 debounce，levenshtein 同步阻塞主线程卡死页面",
     "粘贴超长文本到搜索框",
     "移除无效校验，debounce 内逐条截断 200 字符并提示", "已修复"),
    ("P1", "Web", "components/BookTable.tsx + store/finder.ts",
     "书名行内修改只改内存，导出的 Excel 不含修改，刷新即丢失",
     "改正识别错别字后导出/刷新",
     "修改持久化到 localStorage（按 rid），加载时合并；文案注明导出仍以后端数据为准",
     "部分修复（导出不含前端修改，需后端 PATCH 接口才能彻底解决）"),
    # ---------- P1 未修复 ----------
    ("P1", "后端", "main.py login/register/recognize",
     "登录/注册/识别均无限流无锁定：6 位弱密码可在线爆破；被盗 token 可无限刷视觉 API（每次调用都是费用）",
     "公网部署后被扫描/爆破",
     "IP/账号级失败计数 + 指数退避；识别接口按用户限日额度", "未修复"),
    ("P1", "后端", "main.py require_user",
     "?token= 查询参数所有端点可用；token（30 天有效）会进浏览器历史、uvicorn access log、logs/server.log",
     "任何 <img>/导出链接的使用都会沉淀 token",
     "仅图片端点接受 query token；access log 对 query 脱敏；或改短期一次性签名 URL", "未修复（设计决策待定）"),
    ("P1", "实测", "裁剪交互引导",
     "9 图实测：±0.08 留白不足以覆盖架标——裁剪后 5/9 架标为空、1/9 缺分类号、1/9 读到邻架的错误架标",
     "用户框选只框住书脊不含架标条带",
     "UI 强化引导'框选必须包含书架号标签'；或客户端单独传架标；或架标区域单独再识别一次",
     "未修复（产品交互层，需确认方案）"),
    # ---------- P2 ----------
    ("P2", "后端", "main.py / core/auth.py",
     "用户名枚举双通道：注册 409 直接暴露；登录时用户不存在立即返回、存在则跑 10 万轮 PBKDF2，计时侧信道",
     "攻击者批量探测已注册用户名",
     "注册统一话术；check_user 对不存在用户做 dummy PBKDF2 抹平时延", "未修复"),
    ("P2", "后端", "core/auth.py",
     "token 纯签名串不可吊销：无登出失效、改密不失效、用户删除后仍可用，TTL 30 天",
     "token 泄露后无法止损",
     "verify_token 查库校验用户存在 + token_version；或缩短 TTL + refresh", "未修复"),
    ("P2", "后端", "core/auth.py",
     "PBKDF2 10 万轮低于 OWASP 当前建议（21 万+）", "离线拖库后爆破",
     "提到 21 万以上，哈希串带轮数参数便于迁移", "未修复"),
    ("P2", "后端", "main.py api_recognize",
     "上传图片无大小限制，MAX_EDGE 缩放发生在完整解码之后，1 亿像素'合法'JPEG 解码后占 300MB+ 内存，并发即 OOM",
     "恶意/异常大图上传",
     "Content-Length 预检 + 调低 Image.MAX_IMAGE_PIXELS 到业务合理值", "未修复"),
    ("P2", "后端", "main.py SECRET 初始化",
     "多进程/双实例同启时各自生成 SECRET 互相覆写，token 间歇 401",
     "uvicorn --workers N 或双实例部署",
     "启动加文件锁或从环境变量注入 SECRET", "未修复（单进程部署下不触发）"),
    ("P2", "后端", "main.py _my_records",
     "每次列表/搜索都 glob 全目录 + 逐个读 JSON，无索引无缓存，记录多了延迟线性恶化",
     "单用户几百条记录后",
     "sqlite 建 records 索引表，写入时同步", "未修复"),
    ("P2", "后端", "main.py stream()",
     "客户端中途断连，后台线程池仍把剩余切块全部调完（API 费用照付），结果不落盘",
     "弱网小程序场景常见",
     "识别完成先落盘再推流；或检测断连取消线程池", "未修复（随 P0 任务制一并解决）"),
    ("P2", "后端", "core/search.py",
     "score 结构缺陷：子串匹配恒满分（单字 query'的'匹配一切；OCR 残题反向满分）；置信度权重 0.4 让 sim≥0.167 即过阈值，高置信错书易混入",
     "单字/极短 query；错识别的高置信书名",
     "子串分按长度比缩放；query 最小长度；置信度仅作排序 tiebreak", "未修复"),
    ("P2", "后端", "main.py api_search",
     "query 无服务端长度上限，50 条 × 100KB query 的 SequenceMatcher 单请求占满 worker（CPU DoS）",
     "恶意超长 query（前端已限 200 字符，服务端未限）",
     "服务端 query 限长 64~200 字符", "未修复"),
    ("P2", "小程序", "pages/index/index.js canvas 裁剪",
     "canvas.createImage() 加载原图在部分机型不应用 EXIF 旋转，竖拍 iPhone 照片会裁出错误区域（服务端坐标系一致但客户端裁出的文件内容已错）",
     "竖拍含 EXIF orientation=6 的照片 + canvas 裁剪成功路径",
     "放弃客户端裁剪，统一走'原图 + crop 参数'服务端裁剪（还能省一次本地解码）", "未修复（建议改为服务端裁剪）"),
    ("P2", "小程序", "utils/config.js",
     "BASE_URL 默认 127.0.0.1 真机不可用；真机非调试模式 HTTP 明文被拦截，只会看到'网络错误'",
     "真机预览/体验版",
     "onLaunch 调 /api/health 探测并给明确引导；上线前必须 HTTPS + 合法域名", "未修复（依赖部署决策）"),
    ("P2", "Web", "lib/api.ts recognizeStream",
     "SSE 识别无 AbortController，离开页面后 fetch 继续、服务端持续计费；网络中断已识别进度全部丢失",
     "识别中切页面/断网",
     "支持 AbortSignal + 取消按钮；配合服务端先落盘", "未修复"),
    ("P2", "Web", "store/auth.ts / lib/api.ts",
     "30 天 token 存 localStorage 且拼进图片/导出 URL，进浏览器历史与服务端日志；logout 无法吊销",
     "任一未来 XSS 注入点即可拿走长期 token",
     "缩短有效期；图片改签名 URL；服务端吊销列表", "未修复（与后端 token 策略联动）"),
    # ---------- P3 ----------
    ("P3", "后端", "main.py", "V2 旧记录无 user_id 成孤儿：所有用户不可见、无法导出", "升级前的存量数据", "写一次性迁移脚本补 user_id", "未修复"),
    ("P3", "后端", "main.py", "无删除记录接口，results 目录无限增长；脏记录永久污染搜索", "长期使用", "DELETE /api/results/{rid}（校验归属后删目录）", "未修复"),
    ("P3", "后端", "main.py", "result.json 非原子写，写一半崩溃成坏 JSON 且被静默跳过（无日志）", "进程崩溃", "临时文件 + os.replace；解析失败记 warning", "未修复"),
    ("P3", "后端", "main.py", "RotatingFileHandler 多进程 rotate 在 Windows 上 PermissionError", "--workers N 部署", "文档明确单进程或换 WatchedFileHandler", "未修复"),
    ("P3", "后端", "main.py / core/vision.py", "异常详情回显客户端；/api/health 暴露模型与密钥配置状态；Bearer 大小写敏感；vision.py:69 raise last_err 死代码", "—", "对外通用文案；health 只回 ok；scheme 大小写不敏感；删死代码", "未修复"),
    ("P3", "后端", "core/tiles.py / core/refine.py", "dedup 在裁剪单行场景可能误删并排复本（同名规则）；snap_rows 行聚类链式漂移、docstring 承诺的'消缝'未实现", "同层并排放两本相同的书；透视斜拍", "同名去重附加高度重叠条件；行聚类改直方图锚点", "未修复"),
    ("P3", "小程序", "pages/history 等", "下拉刷新提前结束；handle401 未清 globalData（两处真相源）；搜索 wx:key 用 query 重复告警；SSE 首行 BOM 丢 meta 事件", "—", "逐个微调", "未修复"),
    ("P3", "Web", "多文件", "CameraCapture 流未就绪拍照无反应/无 track.onended；?load= 失败静默；ShelfViewer 定位 nonce 同毫秒碰撞；401 整页跳转丢内存队列；HEIC type 为空被误杀；多标签 logout 不同步", "—", "逐个微调", "未修复"),
    ("P3", "Web", "pages/HistoryPage.tsx", "导出 ids 拼 query 有 URL 超长风险；报错文案是 AxiosError 原始串", "一次导出上百条记录", "ids 改 POST body；格式化错误提示", "未修复"),
]

wb = Workbook()
ws = wb.active
ws.title = "问题清单"
headers = ["级别", "端", "位置", "问题描述", "触发场景", "建议修法", "状态"]
ws.append(headers)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", start_color="2E4B3F")
    c.alignment = Alignment(vertical="center")

LEVEL_FILL = {"P0": "C00000", "P1": "ED7D31", "P2": "FFC000", "P3": "A9D08E"}
for row in ISSUES:
    ws.append(list(row))
    r = ws.max_row
    ws.cell(r, 1).fill = PatternFill("solid", start_color=LEVEL_FILL[row[0]])
    ws.cell(r, 1).font = Font(bold=True, color="FFFFFF" if row[0] in ("P0", "P1") else "000000")
    for col in range(1, 8):
        ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")

for col, w in zip("ABCDEFG", (6, 10, 30, 48, 26, 40, 22)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:G{ws.max_row}"

# 汇总页
ws2 = wb.create_sheet("汇总")
ws2.append(["级别", "总数", "已修复", "部分修复", "未修复"])
for c in ws2[1]:
    c.font = Font(bold=True)
for lv in ("P0", "P1", "P2", "P3"):
    rows = [i for i in ISSUES if i[0] == lv]
    fixed = sum(1 for i in rows if i[6] == "已修复")
    partial = sum(1 for i in rows if i[6].startswith("部分"))
    open_ = len(rows) - fixed - partial
    ws2.append([lv, len(rows), fixed, partial, open_])
ws2.append(["合计", len(ISSUES),
            sum(1 for i in ISSUES if i[6] == "已修复"),
            sum(1 for i in ISSUES if i[6].startswith("部分")),
            sum(1 for i in ISSUES if not i[6].startswith(("已修复", "部分")))])
for col, w in zip("ABCDE", (8, 8, 10, 10, 10)):
    ws2.column_dimensions[col].width = w

out = r"book_finder_v3/tests/V3对抗性检查问题清单.xlsx"
wb.save(out)
print("saved:", out, "rows:", len(ISSUES))
